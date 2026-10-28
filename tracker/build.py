"""Builds a job's spreadsheets from data.json (run inside the job folder):
  jobnumber_PROJECT_FilterLog.xlsx   - the file for the PM, same layout as the approved template (10/28):
                                       tab 1 Filter Log, tab 2 Tech Filter Form (printed), tab 3 Filter Orders
  jobnumber_PROJECT_OpenQuestions.xlsx - just for us, not filed
Unit rows in data.json: [tag, size, qty, status, note, {"1"/"2"/"3"/"CO": "mm/dd/yyyy" date changed}],
the note and dates are optional.
Run with --template to make the blank PM templates in tracker/templates/, --log-only to skip Open Questions."""
import json, os, re, sys
from datetime import datetime
from collections import defaultdict
from itertools import zip_longest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName

NAVY, ORANGE = "2B3C8E", "FC7C36"   # Airco logo colors
HEAD = PatternFill("solid", fgColor=NAVY)
BLANK = PatternFill("solid", fgColor="FFF2CC")   # yellow = still needs info
BAD = PatternFill("solid", fgColor="FFC7CE")     # red = unit not replaced
LABEL = PatternFill("solid", fgColor="DCE0F0")   # light navy
BAND = PatternFill("solid", fgColor="F2F2F2")    # light grey every other row
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

SET_NAMES = {"1": "1st (media roll)", "2": "2nd", "3": "3rd", "Attic": "Attic stock"}
SETS = ["1st (media roll)", "2nd", "3rd", "Attic stock", "RFI", "Change order"]
STATUSES = ["Not started", "Needs filters", "Ordered", "Replaced", "Not replaced", "Washable - not replaced"]
DELIVERY = ["Will call", "Delivered to site", "At Airco warehouse"]
ORDER_CENTER = set(range(1, 14))
ORDER_COLS = ["Order date", "Ordered by", "Filter set or reason", "Vendor", "Vendor SO #", "Our PO #", "MERV",
              "Order total price", "Size", "QTY", "Unit price", "Line total", "Delivery", "Notes"]
ORDER_WIDTHS = [14, 16, 18, 18, 13, 13, 8, 14, 14, 8, 11, 12, 20, 45]
JOB_LEFT = [("Job number", "job_number"), ("Project name", "project"), ("Phase", "phase"), ("Job type", "job_type")]
JOB_RIGHT = [("Airco PM", "airco_pm"), ("Airco APM", "apm"), ("GC", "gc")]   # Ordered by is a column on Filter Orders
MUST_HAVE = {"job_number", "project", "airco_pm"}   # yellow when empty; the rest are optional


def filter_type(note):
    # "Pre-filter, MERV 8." -> Pre, "Final filter ..." -> Final, otherwise left blank
    n = (note or "").lower()
    return "Pre" if n.startswith(("pre-filter", "pre filter")) else "Final" if n.startswith("final filter") else None


def set_name(s):
    return SET_NAMES.get(str(s), s) if s not in (None, "") else None


def header(ws, row, names):
    for c, n in enumerate(names, 1):
        cell = ws.cell(row=row, column=c, value=n)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(wrap_text=True, horizontal="left" if n == "Notes" else "center", vertical="center")
        cell.border = Border(left=thin, right=thin, top=thin, bottom=Side(style="medium", color=ORANGE))
    ws.row_dimensions[row].height = 34


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def dropdown(ws, cells, choices):
    # error off: the list is a shortcut, people can still type something else
    dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv); dv.add(cells)


def title(ws, text, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.cell(row=1, column=1, value=text).font = Font(bold=True, size=16, color=NAVY)
    for c in range(1, ncols + 1):   # orange line under the title
        ws.cell(row=1, column=c).border = Border(bottom=Side(style="medium", color=ORANGE))
    ws.row_dimensions[1].height = 28


def field(ws, row, col, label, last_col, value=None, fill=None, label_last=None):
    # label_last: let a long label ("Airco APM") spread over more than one column
    label_last = label_last or col
    for c in range(col, last_col + 1):
        ws.cell(row=row, column=c).border = BOX
    for c in range(col, label_last + 1):
        ws.cell(row=row, column=c).fill = LABEL
    if label_last > col:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=label_last)
    lab = ws.cell(row=row, column=col, value=label)
    lab.font = Font(bold=True)
    lab.alignment = Alignment(horizontal="left", vertical="center")
    if last_col > label_last + 1:
        ws.merge_cells(start_row=row, start_column=label_last + 1, end_row=row, end_column=last_col)
    v = ws.cell(row=row, column=label_last + 1, value=value)
    v.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    if fill:
        v.fill = fill
    # merged cells do not grow by themselves, so give long values (project name) two lines
    width = sum(ws.column_dimensions[get_column_letter(c)].width or 10 for c in range(label_last + 1, last_col + 1))
    lines = 2 if value and len(str(value)) * 1.1 > width else 1
    ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 0, 20 if lines == 1 else 32)
    return v.coordinate


def job_box(ws, ncols, job=None, copy_from=None, split=None):
    """Job info in two side by side columns (rows 2-5) so the top stays short.
    job: fill in the values. copy_from: (sheet, cells) to show the other tab's values by formula.
    split: (last column of the left side, first and last column of the right side labels)."""
    if split:
        left_last, mid, mid_last = split
    else:
        mid = ncols // 2 + 1
        # right side labels need a wide enough column ("Airco APM"), so skip narrow ones like QTY/MERV
        while (ws.column_dimensions[get_column_letter(mid)].width or 10) < 14 and mid < ncols - 1:
            mid += 1
        left_last, mid_last = mid - 1, mid
    where = {}
    for r, (a, b) in enumerate(zip_longest(JOB_LEFT, JOB_RIGHT), 2):
        for col, lab_last, last, item in [(1, 1, left_last, a), (mid, mid_last, ncols, b)]:
            if not item:
                continue
            lab, key = item
            val = (job or {}).get(key) or None
            fill = BLANK if job is not None and key in MUST_HAVE and not val else None
            where[lab] = field(ws, r, col, lab, last, val, fill, lab_last)
    for c in range(1, ncols + 1):   # orange line between the title and the job box
        cell = ws.cell(row=2, column=c)
        b = cell.border
        cell.border = Border(left=b.left, right=b.right, bottom=b.bottom, top=Side(style="medium", color=ORANGE))
    if copy_from:
        sheet, cells = copy_from
        for lab, ref in where.items():
            src = f"'{sheet}'!{cells[lab]}"
            ws[ref] = f'=IF({src}="","",{src})'
    return where


def table(ws, names, top, nrows, center_cols, wrap_cols):
    """Header at row top, nrows bordered rows under it. Header stays on screen and prints on every page."""
    header(ws, top, names)
    first, last = top + 1, top + nrows
    for r in range(first, last + 1):
        # no fixed height, so rows grow by themselves when a long note wraps
        for c in range(1, len(names) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BOX
            cell.alignment = Alignment(vertical="center", wrap_text=c in wrap_cols,
                                       horizontal="center" if c in center_cols else "left")
    ws.freeze_panes = ws.cell(row=first, column=1)
    ws.print_title_rows = f"{top}:{top}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return first, last


def stripes(ws, first, last, ncols):
    # grey every other row, but never over a yellow/red cell
    for r in range(first, last + 1):
        if r % 2 == 0:
            for c in range(1, ncols + 1):
                cell = ws.cell(row=r, column=c)
                if cell.fill.fill_type is None:
                    cell.fill = BAND


def orders_sheet(ws, log_box, job=None):
    widths(ws, ORDER_WIDTHS)
    title(ws, "FILTER ORDERS", len(ORDER_COLS))
    job_box(ws, len(ORDER_COLS), job, ("Filter Log", log_box))
    ws.row_dimensions[6].height = 8
    return 7


def order_lists(ws, first, last):
    dropdown(ws, f"C{first}:C{last}", SETS)
    dropdown(ws, f"G{first}:G{last}", ["8", "13"])
    dropdown(ws, f"M{first}:M{last}", DELIVERY)


def order_formats(ws, first, last):
    for r in range(first, last + 1):
        ws.cell(row=r, column=1).number_format = "mm/dd/yyyy"
        for c in (8, 11, 12):
            ws.cell(row=r, column=c).number_format = '"$"#,##0.00'


def questions_sheet(ws, job=None):
    title(ws, "OPEN QUESTIONS", 3)
    for r, (lab, key) in enumerate([("Job number", "job_number"), ("Project name", "project"), ("Airco PM", "airco_pm")], 2):
        c = ws.cell(row=r, column=1, value=lab); c.font = Font(bold=True); c.fill = LABEL; c.border = BOX
        v = ws.cell(row=r, column=2, value=(job or {}).get(key) or None); v.border = BOX
        ws.row_dimensions[r].height = 20


# ---------- Blank templates for PMs: python3 build.py --template ----------
# Layout from Angel's PM copy (10/28): Filter Log with a date column per set, a printable Tech Filter Form,
# then Filter Orders. Rows 9-124 line up on the Filter Log and the Tech Filter Form.
TPL_LOG_COLS = ["Equipment Tag", "Filter Type\n(Pre or final)", "Filter size\n(W x L x D)", "QTY", "MERV", "Status",
                "Date Changed\n1st Set", "Date Changed\n2nd Set", "Date Changed\n3rd Set", "Date Changed\nC/O", "Notes"]
TPL_LOG_WIDTHS = [16, 13, 16, 8, 8, 16, 13, 13, 13, 13, 32]
TPL_ROWS = 116
TECH_COLS = ["Equipment Tag", "Filter Type\n(Pre or final)", "Filter size\n(W x L x D)", "QTY", "MERV",
             "Date Changed", "Notes"]
TECH_WIDTHS = [15, 13, 15, 6, 7, 18, 46]   # Date Changed and Notes are wide for handwriting
TECH_ROW_HEIGHT = 26                        # about 1/3 inch per row, room to write
TECH_EXTRA = 10                             # blank rows printed under the units, for units the tech finds
dark = Side(style="thin", color="808080")   # darker lines, light grey fades out on paper
PRINT_BOX = Border(left=dark, right=dark, top=dark, bottom=dark)


def date_only(ws, cells):
    # stops anything that is not a date (error on, unlike the dropdowns)
    dv = DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2099,12,31)",
                        allow_blank=True, showErrorMessage=True, errorStyle="stop", errorTitle="Dates only",
                        error="Type a date, like 10/28/26. Leave it blank if the filters were not changed yet.",
                        showInputMessage=True, promptTitle="Date changed", prompt="Type a date, like 10/28/26")
    ws.add_data_validation(dv); dv.add(cells)
    for row in ws[cells]:
        for cell in row:
            cell.number_format = "mm/dd/yyyy"


def page_numbers(ws, name):
    ws.oddFooter.left.text = name
    ws.oddFooter.right.text = "Page &P of &N"


def tech_sheet(ws, log_box, first, last):
    widths(ws, TECH_WIDTHS)
    ncols = len(TECH_COLS)
    title(ws, "TECH FILTER FORM", ncols)
    job_box(ws, ncols, copy_from=("Filter Log", log_box), split=(4, 5, 6))
    ws.row_dimensions[3].height = 30
    # filled in by the tech
    field(ws, 5, 5, "Tech name", ncols, label_last=6)
    v = field(ws, 6, 1, "Filter set\n(circle one)", ncols)
    ws[v] = "1st           2nd           3rd           C/O"
    ws[v].font = Font(bold=True, size=12)
    ws["A6"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[6].height = 32
    ws.merge_cells(start_row=7, start_column=1, end_row=7, end_column=ncols)
    tip = ws.cell(row=7, column=1, value="Check each unit as you go. If a tag, size or QTY is wrong, cross it out and write "
                  "the right one. Write the date you changed the filters and any notes. Add missing units on the blank rows.")
    tip.font = Font(italic=True, size=10, color="404040")
    tip.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[7].height = 30
    table(ws, TECH_COLS, 8, last - first + 1, center_cols=set(range(1, 7)), wrap_cols={2, 7})
    ws.row_dimensions[8].height = 34
    for r in range(first, last + 1):
        ws.row_dimensions[r].height = TECH_ROW_HEIGHT
        for c in range(1, ncols + 1):
            ws.cell(row=r, column=c).border = PRINT_BOX
        for c in range(1, 6):   # unit info from the Filter Log, blank (not 0) when the log row is empty
            src = f"'Filter Log'!{get_column_letter(c)}{r}"
            ws.cell(row=r, column=c, value=f'=IF({src}="","",{src})')
    # print setup: portrait for a clipboard, 1 page wide, header row on every page
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top, ws.page_margins.bottom = 0.5, 0.6
    ws.page_margins.header, ws.page_margins.footer = 0.3, 0.3
    ws.print_options.horizontalCentered = True
    page_numbers(ws, "Tech Filter Form")
    # print only the rows the Filter Log uses, plus TECH_EXTRA blank rows (Excel works this out when printing)
    log = lambda col: f"COUNTA('Filter Log'!${col}${first}:${col}${last})"
    rows = f"MIN({last - first + 1},MAX({log('A')},{log('C')})+{TECH_EXTRA})"
    ws.defined_names["_xlnm.Print_Area"] = DefinedName(
        "_xlnm.Print_Area", attr_text=f"'{ws.title}'!$A$1:INDEX('{ws.title}'!$G${first}:$G${last},{rows})")


def new_log_sheet(ws, nrows, job=None):
    """Filter Log tab in the approved PM layout (10/28). Header row 8, units from row 9."""
    widths(ws, TPL_LOG_WIDTHS)
    ncols = len(TPL_LOG_COLS)
    title(ws, "FILTER LOG", ncols)
    log_box = job_box(ws, ncols, job, split=(4, 7, 8))
    ws.row_dimensions[3].height = 30   # room for a long project name on two lines
    ws.row_dimensions[6].height = ws.row_dimensions[7].height = 8
    note = re.sub(r"\s*Yellow = still needs info\.?", "", (job or {}).get("note") or "").strip()
    if note:   # optional job note on row 6, the table still starts on row 8
        ws.merge_cells(start_row=6, start_column=1, end_row=6, end_column=ncols)
        c = ws.cell(row=6, column=1, value=note)
        c.font = Font(italic=True, color="404040"); c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[6].height = 15 * (len(note) // 180 + 1) + 4
    first, last = table(ws, TPL_LOG_COLS, 8, nrows, center_cols=set(range(1, 11)), wrap_cols={2, 11})
    dropdown(ws, f"F{first}:F{last}", STATUSES)
    date_only(ws, f"G{first}:J{last}")
    page_numbers(ws, "Filter Log")
    return log_box, first, last


def build_templates():
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "templates")
    os.makedirs(out, exist_ok=True)

    def bands(ws, rng):
        # light grey every other row; added last so the red/yellow rules win
        ws.conditional_formatting.add(rng, FormulaRule(formula=["MOD(ROW(),2)=0"], fill=BAND))

    # 1. Filter log (main tab). PMs fill this in from submittals; one date column per set.
    wb = Workbook(); ws = wb.active; ws.title = "Filter Log"
    log_box, first, last = new_log_sheet(ws, TPL_ROWS)
    ws.conditional_formatting.add(f"A{first}:K{last}", FormulaRule(formula=[f'$F{first}="Not replaced"'], fill=BAD, stopIfTrue=True))
    bands(ws, f"A{first}:K{last}")

    # 2. Tech filter form: printed and handed to the tech. Unit info comes from the Filter Log by formula,
    # the tech writes the date and notes by hand.
    tech_sheet(wb.create_sheet("Tech Filter Form"), log_box, first, last)

    # 3. Filter orders: last tab. Only a purchase log, unit/size work lives in the Filter Log.
    ws = wb.create_sheet("Filter Orders")
    top = orders_sheet(ws, log_box)
    first, last = table(ws, ORDER_COLS, top, 60, center_cols=ORDER_CENTER, wrap_cols={2, 4, 13, 14})
    order_formats(ws, first, last)
    for r in range(first, last + 1):
        ws.cell(row=r, column=12, value=f'=IF(OR(J{r}="",K{r}=""),"",J{r}*K{r})')
    order_lists(ws, first, last)
    bands(ws, f"A{first}:N{last}")
    widths(ws, ORDER_WIDTHS)
    page_numbers(ws, "Filter Orders")
    wb.save(os.path.join(out, "TEMPLATE_FilterLog.xlsx"))

    # 4. Open questions (just for us, not filed with the PM)
    wb = Workbook(); ws = wb.active; ws.title = "Open Questions"
    questions_sheet(ws)
    ws.cell(row=5, column=2, value='=COUNTIF(A8:A107,"Open")&" open - ready to send when this is 0 and every unit and filter is accounted for"').font = Font(bold=True)
    first, last = table(ws, ["Status", "Question", "Who knows"], 7, 100, center_cols={1}, wrap_cols={2, 3})
    dropdown(ws, f"A{first}:A{last}", ["Open", "Answered", "Later"])
    ws.conditional_formatting.add(f"A{first}:A{last}", FormulaRule(formula=[f'$A{first}="Open"'], fill=BLANK, stopIfTrue=True))
    bands(ws, f"A{first}:C{last}")
    widths(ws, [14, 90, 22])
    wb.save(os.path.join(out, "TEMPLATE_OpenQuestions.xlsx"))
    print("Templates saved in", out)


if "--template" in sys.argv:
    build_templates()
    sys.exit()

data = json.load(open("data.json"))
job = data["job"]
# file names: jobnumber_PROJECTNAME_what-it-is, no spaces (e.g. 2250473_AISDLANGFORD_FilterLog.xlsx)
prefix = f"{job['job_number']}_{job.get('short_name') or re.sub(r'[^A-Za-z0-9]', '', job['project']).upper()}_"
EXTRA = 20   # empty rows under the data, with the same dropdowns, for adding by hand

# ---------- 1. Filter log (tab 1 of the PM file) ----------
wb = Workbook(); ws = wb.active; ws.title = "Filter Log"
n = len(data["units"]) or data.get("blank_rows", 0)
log_box, first, last = new_log_sheet(ws, n + EXTRA, job)
DATE_COL = {"1": 7, "2": 8, "3": 9, "CO": 10}   # Date Changed 1st / 2nd / 3rd / C/O
for i, (unit, size, qty, status, *more) in enumerate(data["units"], first):
    done = status == "Replaced"
    note = more[0] if more else None
    dates = more[1] if len(more) > 1 else {}   # optional 6th item: {"2": "10/28/2026", ...}
    vals = [unit, filter_type(note), size, qty, 8 if done else None, status, None, None, None, None, note]
    for c, v in enumerate(vals, 1):
        ws.cell(row=i, column=c, value=v)
    for k, d in (dates or {}).items():
        ws.cell(row=i, column=DATE_COL[k], value=datetime.strptime(d, "%m/%d/%Y"))
    # unknown unit tag (size tally, no unit list yet), size or qty: mark yellow
    if str(unit).startswith("TBD"):
        ws.cell(row=i, column=1).fill = BLANK
    if "?" in size:
        ws.cell(row=i, column=3).fill = BLANK
    if qty is None:
        ws.cell(row=i, column=4).fill = BLANK
    if status == "Not replaced":
        for c in range(1, len(TPL_LOG_COLS) + 1):
            ws.cell(row=i, column=c).fill = BAD
    elif done and not dates:
        ws.cell(row=i, column=DATE_COL["2"]).fill = BLANK   # replaced (2nd set) but no date yet
    elif not done and not status.startswith("Washable"):
        ws.cell(row=i, column=5).fill = BLANK   # MERV
# no unit list yet: leave blank yellow rows to fill in
if not data["units"]:
    for i in range(first, first + n):
        for c in range(1, len(TPL_LOG_COLS) + 1):
            ws.cell(row=i, column=c).fill = BLANK
stripes(ws, first, last, len(TPL_LOG_COLS))

# ---------- 2. Tech filter form (tab 2, printed for the tech) ----------
tech_sheet(wb.create_sheet("Tech Filter Form"), log_box, first, last)

# ---------- 3. Filter orders (tab 3 of the PM file) ----------
ws = wb.create_sheet("Filter Orders")
top = orders_sheet(ws, log_box, job)
n = sum(len(o["lines"]) for o in data["orders"])
first, last = table(ws, ORDER_COLS, top, n + EXTRA, center_cols=ORDER_CENTER, wrap_cols={2, 4, 13, 14})
order_formats(ws, first, last)
r = first
for o in data["orders"]:
    for k, (size, qty, price) in enumerate(o["lines"]):
        first_line = k == 0
        # order-wide info (total, notes) only on the first line, so a big order is easy to read
        note = o["note"] if first_line else None
        if first_line and o.get("units"):
            note = (note + " " if note else "") + "Units: " + o["units"]
        # who placed this order; falls back to the job's "ordered_by" only for old data without it
        vals = [o["date"], o.get("ordered_by") or job.get("ordered_by") or None, set_name(o["set"]), o["vendor"], o["vendor_so"], o["po_number"] or None, o["merv"],
                o["price"] if first_line else None, size, qty, price, f"=J{r}*K{r}", o["delivery"], note]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if v is None and (c in (2, 3, 6) or (first_line and c in (1, 4, 5, 7, 8, 13))):
                cell.fill = BLANK
        r += 1
for rr in range(r, last + 1):
    ws.cell(row=rr, column=12, value=f'=IF(OR(J{rr}="",K{rr}=""),"",J{rr}*K{rr})')
stripes(ws, first, last, len(ORDER_COLS))
order_lists(ws, first, last)
widths(ws, ORDER_WIDTHS)
page_numbers(ws, "Filter Orders")
wb.save(prefix + "FilterLog.xlsx")

# console summary: ordered vs needed per size (for us; no longer a sheet in the PM file)
ordered = defaultdict(int); need = defaultdict(int)
for o in data["orders"]:
    for s, q, _ in o["lines"]: ordered[s] += q
for _, s, q, st, *_ in data["units"]:
    if st == "Replaced": need[s] += q or 0
for s in sorted(set(ordered) | set(need)):
    print(f"{s:10} ordered {ordered[s]:3}  needed {need[s]:3}  diff {ordered[s]-need[s]:+d}")

# ---------- 4. Open questions (just for us, not filed with the PM) ----------
if "--log-only" in sys.argv:   # only rebuild the PM file
    sys.exit()
wb = Workbook(); ws = wb.active; ws.title = "Open Questions"
questions_sheet(ws, job)
n_open = sum(1 for q in data["questions"] if q[0] == "Open")
ws.cell(row=5, column=2, value=f"{n_open} open - ready to send when this is 0 and every unit and filter is accounted for").font = Font(bold=True)
first, last = table(ws, ["Status", "Question", "Who knows"], 7, len(data["questions"]) + EXTRA, center_cols={1}, wrap_cols={2, 3})
for i, (st, q, who) in enumerate(data["questions"], first):
    for c, v in enumerate([st, q, who], 1):
        ws.cell(row=i, column=c, value=v)
    if st == "Open":
        ws.cell(row=i, column=1).fill = BLANK
stripes(ws, first, last, 3)
dropdown(ws, f"A{first}:A{last}", ["Open", "Answered", "Later"])
widths(ws, [14, 90, 22])
wb.save(prefix + "OpenQuestions.xlsx")
